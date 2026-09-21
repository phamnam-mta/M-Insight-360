import docx

from .types import ExtractedDocument, ExtractedTable


def extract_docx(file_path: str, filename: str) -> ExtractedDocument:
    d = docx.Document(file_path)
    text_parts = [p.text for p in d.paragraphs if p.text.strip()]
    tables = []
    for i, t in enumerate(d.tables):
        rows = [[cell.text for cell in row.cells] for row in t.rows]
        tables.append(ExtractedTable(rows=rows, sheet_or_page=f"table {i + 1}"))
    return ExtractedDocument(
        filename=filename,
        doc_type="docx",
        text="\n".join(text_parts),
        tables=tables,
        extraction_method="docx",
        confidence=1.0,
    )
