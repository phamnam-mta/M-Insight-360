import csv

import openpyxl

from .types import ExtractedDocument, ExtractedTable


def extract_xlsx(file_path: str, filename: str) -> ExtractedDocument:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    tables = []
    text_parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            str_row = ["" if cell is None else str(cell) for cell in row]
            rows.append(str_row)
            text_parts.append(" | ".join(str_row))
        tables.append(ExtractedTable(rows=rows, sheet_or_page=sheet_name))
    return ExtractedDocument(
        filename=filename,
        doc_type="xlsx",
        text="\n".join(text_parts),
        tables=tables,
        extraction_method="spreadsheet",
        confidence=1.0,
    )


def extract_csv(file_path: str, filename: str) -> ExtractedDocument:
    with open(file_path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    text = "\n".join(" | ".join(row) for row in rows)
    doc_type = "csv"
    return ExtractedDocument(
        filename=filename,
        doc_type=doc_type,
        text=text,
        tables=[ExtractedTable(rows=rows, sheet_or_page="csv")],
        extraction_method="spreadsheet",
        confidence=1.0,
    )
