import io

import docx

_TEMPLATE_PATH = "docs/templates/MB01A_QT.RR.038_lan_3.docx"
_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def write_value_after_label(table, row_idx: int, cell_idx: int, value) -> None:
    if not value:
        return
    cell = table.rows[row_idx].cells[cell_idx]
    cell.paragraphs[0].add_run(str(value))


def check_box(table, row_idx: int, cell_idx: int, occurrence: int = 0) -> None:
    cell = table.rows[row_idx].cells[cell_idx]
    checkboxes = cell._tc.findall(f".//{_W_NS}checkBox")
    box = checkboxes[occurrence]
    default = box.find(f"{_W_NS}default")
    default.set(f"{_W_NS}val", "1")


_GENDER_ROW = 4
_MARITAL_ROW = 14
# Unlike the gender checkboxes (row 4), the 4 marital-status boxes each live
# in their OWN cell (verified: cells 9/26/40/53 each contain exactly one
# w:checkBox) rather than 4 occurrences within one cell — map to cell index,
# occurrence is always 0.
_MARITAL_CELL = {"single": 9, "married": 26, "divorced": 40, "widowed": 53}


def build_mb01a_docx(case: dict) -> bytes:
    document = docx.Document(_TEMPLATE_PATH)
    table = document.tables[0]

    customer = case.get("customer") or {}
    write_value_after_label(table, 3, 0, customer.get("full_name"))
    if customer.get("gender") == "male":
        check_box(table, _GENDER_ROW, 0, occurrence=0)
    elif customer.get("gender") == "female":
        check_box(table, _GENDER_ROW, 0, occurrence=1)
    write_value_after_label(table, 5, 0, customer.get("id_number"))
    write_value_after_label(table, 9, 0, customer.get("permanent_address"))
    write_value_after_label(table, 10, 0, customer.get("temporary_address"))
    write_value_after_label(table, 11, 0, customer.get("contact_address"))
    write_value_after_label(table, 13, 0, customer.get("email"))
    marital_cell = _MARITAL_CELL.get(customer.get("marital_status"))
    if marital_cell is not None:
        check_box(table, _MARITAL_ROW, marital_cell, occurrence=0)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
