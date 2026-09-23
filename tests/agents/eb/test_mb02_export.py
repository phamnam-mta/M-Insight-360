import zipfile

import docx

from app.agents.eb.mb02_export import (
    _MISSING_PLACEHOLDER, _format_number_vn, build_mb02_docx,
)


def test_format_number_vn_uses_vietnamese_separators():
    assert _format_number_vn(128_061_897.6, 0) == "128.061.898"
    assert _format_number_vn(0.14, 2) == "0,14"


def test_format_number_vn_none_is_missing_placeholder():
    assert _format_number_vn(None, 0) == _MISSING_PLACEHOLDER


def test_build_mb02_docx_returns_valid_docx_bytes_and_fill_log():
    computed = {
        "customer_profile": {"customer_name": "ACME", "tax_id": "0101234567"},
        "financial_inputs": {}, "credit_engine": {}, "risk_flags": [],
        "export_gate": {"verdict": "XUAT", "block_type": None, "signal_count": 0, "signals": [], "data_warnings": [], "reasons": []},
    }
    docx_bytes, fill_log = build_mb02_docx(computed)
    assert docx_bytes[:2] == b"PK"  # a .docx is a zip archive
    with zipfile.ZipFile(__import__("io").BytesIO(docx_bytes)) as zf:
        assert "word/document.xml" in zf.namelist()
    assert "verdict" in fill_log
    assert isinstance(fill_log["cells"], list)


def _make_computed(**financial_inputs_overrides):
    financial_inputs = {
        "net_revenue_vnd": 128_061_897_765.0, "pbt_vnd": 12_160_680_287.0,
        "pat_vnd": 9_000_000_000.0, "interest_expense_vnd": 12_131_595_580.0,
        "equity_vnd": 100_000_000_000.0, "current_assets_vnd": 200_000_000_000.0,
        "non_current_assets_vnd": 50_000_000_000.0, "short_term_debt_vnd": 78_810_636_239.0,
        "long_term_debt_vnd": 0.0, "charter_capital_vnd": 50_000_000_000.0,
        **financial_inputs_overrides,
    }
    return {
        "customer_profile": {"customer_name": "CÔNG TY ACME", "tax_id": "0101234567"},
        "financial_inputs": financial_inputs,
        "credit_engine": {
            "current_ratio": {"value": 2.5, "status": "OK"},
            "total_borrowings": {"value": 78_810_636_239.0, "status": "OK"},
        },
        "risk_flags": [],
        "export_gate": {"verdict": "XUAT", "block_type": None, "signal_count": 0, "signals": [], "data_warnings": [], "reasons": []},
    }


def _reload(docx_bytes):
    return docx.Document(__import__("io").BytesIO(docx_bytes))


def test_s2_1_fills_customer_name():
    docx_bytes, _ = build_mb02_docx(_make_computed())
    doc = _reload(docx_bytes)
    row = next(r for r in doc.tables[2].rows if "Tên Doanh nghiệp" in r.cells[0].text)
    assert "CÔNG TY ACME" in row.cells[1].text


def test_s2_2_fills_revenue_in_nam_n_column():
    docx_bytes, _ = build_mb02_docx(_make_computed())
    doc = _reload(docx_bytes)
    table31 = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == "Chỉ tiêu" and len(t.rows) == 14)
    row = next(r for r in table31.rows if "Tổng doanh thu" in r.cells[0].text)
    assert row.cells[3].text.strip() == "128.062"  # triệu đồng, rounded
    assert row.cells[1].text.strip() == "[Chưa xác định từ hồ sơ tải lên]"  # Năm N-2 never borrowed


def test_s2_2_leverage_ratio_two_decimals():
    docx_bytes, _ = build_mb02_docx(_make_computed())
    doc = _reload(docx_bytes)
    table31 = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == "Chỉ tiêu" and len(t.rows) == 14)
    row = next(r for r in table31.rows if "Hệ số đòn bẩy" in r.cells[0].text)
    assert row.cells[3].text.strip() == "0,79"  # 78.8bn / 100bn


def test_s2_2_missing_field_is_placeholder_not_zero():
    computed = _make_computed(pat_vnd=None)
    docx_bytes, _ = build_mb02_docx(computed)
    doc = _reload(docx_bytes)
    table31 = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == "Chỉ tiêu" and len(t.rows) == 14)
    row = next(r for r in table31.rows if "Lợi nhuận sau thuế" in r.cells[0].text)
    assert row.cells[3].text.strip() == "[Chưa xác định từ hồ sơ tải lên]"


def test_s2_3_fills_long_term_capital_breakdown():
    computed = _make_computed()
    computed["credit_engine"]["long_term_capital"] = {"value": 100_000_000_000.0, "status": "OK"}
    computed["credit_engine"]["nwc"] = {"value": 150_000_000_000.0, "status": "OK"}
    docx_bytes, _ = build_mb02_docx(computed)
    doc = _reload(docx_bytes)
    table32 = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == "Chỉ tiêu" and len(t.rows) == 6)
    row = next(r for r in table32.rows if r.cells[0].text.strip() == "1. Nguồn vốn dài hạn")
    assert row.cells[3].text.strip() == "100.000"
    row = next(r for r in table32.rows if r.cells[0].text.strip() == "3. Vốn lưu động thường xuyên")
    assert row.cells[3].text.strip() == "150.000"


def test_s2_3_replaces_financial_health_commentary_paragraph():
    computed = _make_computed()
    computed["credit_engine"]["long_term_capital"] = {"value": 100_000_000_000.0, "status": "OK"}
    computed["credit_engine"]["nwc"] = {"value": 150_000_000_000.0, "status": "OK"}
    computed["why"] = ["Doanh thu tăng trưởng ổn định.", "Đòn bẩy ở mức an toàn."]
    docx_bytes, _ = build_mb02_docx(computed)
    doc = _reload(docx_bytes)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Doanh thu tăng trưởng ổn định." in text
    assert "Khuyến nghị sơ bộ từ dữ liệu BCTC" in text


def test_s2_5_fills_nam_0_only_never_plan_years():
    docx_bytes, _ = build_mb02_docx(_make_computed())
    doc = _reload(docx_bytes)
    table37 = next(
        t for t in doc.tables
        if len(t.rows) == 20 and len(t.columns) >= 2 and t.rows[0].cells[1].text.strip() == "Chỉ tiêu"
    )
    row = next(r for r in table37.rows if "Tổng doanh thu" in r.cells[1].text)
    assert row.cells[2].text.strip() == "128.062"  # Năm 0
    assert row.cells[3].text.strip() == ""  # Năm 1 — plan year, never written
    assert row.cells[4].text.strip() == ""  # Năm 2 — plan year, never written


def test_s2_4_fills_reference_limit_paragraph():
    computed = _make_computed(receivables_vnd=1_030_523_666.0)
    docx_bytes, _ = build_mb02_docx(computed)
    doc = _reload(docx_bytes)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Hạn mức tài trợ theo phương án đầu ra dự kiến: 824" in text  # 1.030.523.666 * 80% / 1e6 ≈ 824 triệu
    assert "Số dư phải thu cuối kỳ" in text


def test_s2_4_missing_receivables_leaves_placeholder_note():
    computed = _make_computed(receivables_vnd=None)
    docx_bytes, _ = build_mb02_docx(computed)
    doc = _reload(docx_bytes)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Hạn mức tài trợ theo phương án đầu ra dự kiến: [Chưa xác định" in text


def test_s8_legal_checkbox_row_never_written():
    # "Tình trạng khách hàng" (row 5 of table 2) is an S1.3 checkbox row —
    # must remain byte-for-byte the template's own text.
    docx_bytes, _ = build_mb02_docx(_make_computed())
    doc = _reload(docx_bytes)
    row = next(r for r in doc.tables[2].rows if r.cells[0].text.strip() == "Tình trạng khách hàng")
    assert row.cells[1].text.strip() == "KH mới       KH hiện hữu"


def test_s8_cif_field_keeps_template_ellipsis():
    # CIF has no BCTC source at all (S8) — must never become
    # "[Chưa xác định từ hồ sơ tải lên]", only the template's own "……".
    docx_bytes, _ = build_mb02_docx(_make_computed())
    doc = _reload(docx_bytes)
    row = next(r for r in doc.tables[2].rows if "Đăng ký kinh doanh" == r.cells[0].text.strip())
    assert "……" in row.cells[1].text
    assert "[Chưa xác định" not in row.cells[1].text
