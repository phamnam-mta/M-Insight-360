import io

import docx

from app.agents.rb_portal.mb01a_export import build_mb01a_docx, check_box, write_value_after_label

TEMPLATE_PATH = "docs/templates/MB01A_QT.RR.038_lan_3.docx"


def test_write_value_after_label_appends_to_the_cell():
    document = docx.Document(TEMPLATE_PATH)
    table = document.tables[0]
    write_value_after_label(table, 3, 0, "NGUYEN VAN A")
    assert table.rows[3].cells[0].text == "Họ tên: NGUYEN VAN A"


def test_write_value_after_label_is_noop_for_falsy_value():
    document = docx.Document(TEMPLATE_PATH)
    table = document.tables[0]
    original = table.rows[3].cells[0].text
    write_value_after_label(table, 3, 0, None)
    assert table.rows[3].cells[0].text == original


def test_check_box_sets_default_to_1_for_the_targeted_occurrence_only():
    document = docx.Document(TEMPLATE_PATH)
    table = document.tables[0]
    check_box(table, 4, 0, occurrence=0)  # "Nam"
    xml = table.rows[4].cells[0]._tc.xml
    # exactly one w:default now reads val="1" within this cell (Nữ stays "0")
    assert xml.count('w:default w:val="1"') == 1
    assert xml.count('w:default w:val="0"') == 1


def test_build_mb01a_docx_returns_openable_bytes_with_customer_name():
    case = {
        "customer": {"full_name": "NGUYEN VAN A", "gender": "male", "id_number": "001234567890"},
        "legal": None, "income": None, "loan": None, "collateral": None, "other": None,
    }
    result = build_mb01a_docx(case)
    assert isinstance(result, bytes)
    document = docx.Document(io.BytesIO(result))
    assert "NGUYEN VAN A" in document.tables[0].rows[3].cells[0].text


def test_build_mb01a_docx_leaves_fields_blank_for_empty_case():
    case = {"customer": None, "legal": None, "income": None, "loan": None, "collateral": None, "other": None}
    result = build_mb01a_docx(case)
    document = docx.Document(io.BytesIO(result))
    assert document.tables[0].rows[3].cells[0].text == "Họ tên: "
